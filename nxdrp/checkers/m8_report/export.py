"""可选 Word/PDF 导出。缺失依赖时明确报错，不写伪文件。"""
from pathlib import Path
from xml.sax.saxutils import escape
from .core import fmt, SEVERITY
from .render import details


def export_docx(report, directory):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        # Word 默认模板的标题可能携带段落底线，显式移除。
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    for style in doc.styles:
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(12)
    # 真实页码字段，Word 打开时可更新。
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("设计评审报告  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if widths:
            for col, width in zip(t.columns, widths):
                col.width = Inches(width)
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = str(text)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            for j, cell in enumerate(row.cells):
                if widths:
                    cell.width = Inches(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "DDE8EE" if i == 0 else ("F5F7F9" if i%2 == 0 else "FFFFFF"))
                props.append(shade)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:"+side)
                    el.set(qn("w:w"), "90")
                    el.set(qn("w:type"), "dxa")
                    margins.append(el)
                props.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    if i == 0:
                        p.paragraph_format.keep_with_next = True
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.bold = i == 0
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:"+side)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "D9D9D9")
            borders.append(el)
        t._tbl.tblPr.append(borders)
        return t

    doc.add_heading(report["title"], 0)
    doc.add_paragraph(f"模型 {report['model_id']}\n生成时间 {report['created_at']}\n归档编号 {report['report_id']}")
    s = report["summary"]
    doc.add_paragraph(f"本报告列出全部 {s['catalog_count']} 个检查项，成功执行 {s['executed_count']} 项，记录 {s['finding_count']} 条问题。"
                      f"报告证据存在 {s['quality_issue_count']} 处待补充或复核。")
    acceptance = report['acceptance']
    doc.add_paragraph(f"报告状态  {acceptance['status']}。问题四要素完整 {acceptance['four_fields_complete_count']}/{acceptance['finding_count']} 条。数据完整性检查不代表产品合格。")
    if report["is_mock"]:
        doc.add_paragraph("本报告使用开发示例数据，用于验证报告流程，不代表产品检测结论。")
    doc.add_paragraph("偏差量按实测值减阈值展示，正负号不单独代表是否合格；检测结论沿用上游结果。未提供执行记录的项目列为未执行。")
    for warning in report["warnings"]:
        doc.add_paragraph("待复核  "+warning)
    doc.add_heading("四十四项结果索引", 1)
    index_table = table(["检查项", "名称", "状态", "数量", "结果位置"],
          [[x["id"], x["name"], x["status_label"], x["count"], "§"+x["id"]] for x in report["sections"]],
          [.8, 2.0, 1.45, .45, .8])
    for row, item in zip(index_table.rows[1:], report['sections']):
        cell = row.cells[-1]
        cell.text = ''
        link = OxmlElement('w:hyperlink')
        link.set(qn('w:anchor'), item['anchor'].replace('-', '_'))
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = '§'+item['id']
        run.append(text)
        link.append(run)
        cell.paragraphs[0]._p.append(link)
    doc.add_heading("问题详情", 1)
    for index, item in enumerate(report["sections"]):
        paragraph = doc.add_heading("§"+item["id"]+" "+item["name"], 2)
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(1000+index))
        start.set(qn('w:name'), item['anchor'].replace('-', '_'))
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(1000+index))
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)
        doc.add_paragraph(item["status_label"]+("；"+item["reason"] if item["reason"] else ""))
        for f in item["findings"]:
            doc.add_heading(f["finding_id"]+" "+f["severity"]+" "+SEVERITY.get(f["severity"], "待补充"), 3)
            doc.add_paragraph(str(f["source"].get("message", "")))
            if f["image"]:
                doc.add_picture(str(Path(directory)/f["image"]), width=Inches(5.7))
                doc.add_paragraph(f["image_kind"])
            # 长建议用正文，紧凑字段用表格。
            table(["字段", "结果"], [(k,v) for k,v in details(f) if k not in ("整改建议", "证据说明")], [1.1, 4.4])
            doc.add_paragraph("整改建议  "+f["suggestion"])
            if f["source"]["evidence"].get("note"):
                doc.add_paragraph("证据说明  "+f["source"]["evidence"]["note"])
            for q in f["quality"]:
                doc.add_paragraph("待补充或复核  "+q)
            for note in f.get('advisories', []):
                doc.add_paragraph('人工确认  '+note)
    doc.add_heading("追溯信息", 1)
    doc.add_paragraph(f"模型来源 {report['model_source']}\n坐标单位 {report['model_units']}\n建议库版本 {report['knowledge_version']}\n结果 SHA256 {report['source_digest']}")
    doc.save(str(Path(directory)/"report.docx"))


def export_pdf(report, directory):
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
    font_name = "M8Chinese"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        font_path = Path("C:/Windows/Fonts/msyh.ttc")
        if not font_path.exists():
            font_path = Path("C:/Windows/Fonts/simhei.ttf")
        if font_path.exists():
            pdfmetrics.registerFont(TTFont(font_name, str(font_path), subfontIndex=0))
        else:
            font_name = "STSong-Light"
            pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    normal = ParagraphStyle("m8", fontName=font_name, fontSize=9.5, leading=15, spaceAfter=7, wordWrap="CJK")
    small = ParagraphStyle("small", parent=normal, fontSize=8, leading=12, spaceAfter=0)
    heading = ParagraphStyle("h", parent=normal, fontSize=16, leading=23, spaceBefore=17, spaceAfter=10)
    sub = ParagraphStyle("sub", parent=normal, fontSize=11.5, leading=18, spaceBefore=10, keepWithNext=True)
    title = ParagraphStyle("title", parent=normal, fontSize=25, leading=34, spaceAfter=18)
    p = lambda text, style=normal: Paragraph(escape(str(text)).replace("\n", "<br/>"), style)
    def table(headers, rows, widths):
        data = [[p(x, small) for x in headers]]+[[x if isinstance(x, Paragraph) else p(x, small) for x in row] for row in rows]
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#dde8ee")),
                               ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f5f7f9")]),
                               ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#d9d9d9")),
                               ("VALIGN", (0,0), (-1,-1), "TOP"), ("LEFTPADDING",(0,0),(-1,-1),7),
                               ("RIGHTPADDING",(0,0),(-1,-1),7), ("TOPPADDING",(0,0),(-1,-1),6),
                               ("BOTTOMPADDING",(0,0),(-1,-1),6)]))
        return t
    class Bookmark(Paragraph):
        def __init__(self, text, key):
            super().__init__(escape(text), sub)
            self.key = key
        def draw(self):
            self.canv.bookmarkHorizontal(self.key, 0, self.height)
            self.canv.addOutlineEntry(self.getPlainText(), self.key, 0, False)
            super().draw()
    elements = [p(report["title"], title), p(f"模型 {report['model_id']}\n生成时间 {report['created_at']}\n归档编号 {report['report_id']}")]
    s = report["summary"]
    acceptance = report['acceptance']
    elements.append(p(f"报告状态  {acceptance['status']}。问题四要素完整 {acceptance['four_fields_complete_count']}/{acceptance['finding_count']} 条。数据完整性检查不代表产品合格。"))
    elements += [p(f"全部 {s['catalog_count']} 项均有结果位置；成功执行 {s['executed_count']} 项，问题 {s['finding_count']} 条，证据待补充 {s['quality_issue_count']} 处。"),
                 p("开发示例数据，用于验证报告流程，不代表产品检测结论。" if report["is_mock"] else "请结合执行状态与待补充证据开展人工评审。"),
                 p("偏差按实测值减阈值展示，正负号不单独代表合格与否。无执行记录的项目列为未执行。")]
    for warning in report["warnings"]:
        elements.append(p("待复核  "+warning))
    elements.append(p("四十四项结果索引", heading))
    rows = []
    for item in report["sections"]:
        link = Paragraph(f'<link href="#{item["anchor"]}">§{item["id"]}</link>', small)
        rows.append([item["id"], item["name"], item["status_label"], str(item["count"]), link])
    elements += [table(["检查项", "名称", "状态", "数量", "结果位置"], rows, [60, 155, 100, 35, 75]), PageBreak(), p("问题详情", heading)]
    for item in report["sections"]:
        elements += [Bookmark("§"+item["id"]+" "+item["name"], item["anchor"]), p(item["status_label"]+("；"+item["reason"] if item["reason"] else ""))]
        for f in item["findings"]:
            elements += [p(f["finding_id"]+" "+f["severity"]+" "+SEVERITY.get(f["severity"], "待补充"), sub), p(f["source"].get("message", ""))]
            if f["image"]:
                from PIL import Image as PILImage
                path = Path(directory)/f["image"]
                with PILImage.open(path) as im:
                    w, h = im.size
                scale = min(425/w, 220/h)
                elements += [Image(str(path), w*scale, h*scale, hAlign="LEFT"), p(f["image_kind"], small), Spacer(1, 8)]
            elements.append(table(["字段", "结果"], [(k,v) for k,v in details(f) if k not in ("整改建议", "证据说明")], [75, 350]))
            elements += [Spacer(1, 8), p("整改建议  "+f["suggestion"])]
            for q in f["quality"]:
                elements.append(p("待补充或复核  "+q))
            for note in f.get('advisories', []):
                elements.append(p('人工确认  '+note))
    elements += [p("追溯信息", heading), p(f"模型来源 {report['model_source']}\n坐标单位 {report['model_units']}\n建议库版本 {report['knowledge_version']}\n结果 SHA256 {report['source_digest']}")]
    def footer(c, doc):
        c.setFont(font_name, 8)
        c.setFillColor(colors.HexColor("#687d8d"))
        c.drawString(46, 25, "连接器设计检查报告")
        c.drawRightString(565, 25, str(doc.page))
    doc = SimpleDocTemplate(str(Path(directory)/"report.pdf"), pagesize=(612,792),
                            leftMargin=46, rightMargin=46, topMargin=42, bottomMargin=42,
                            title=report["title"], author="M8 报告与知识库")
    doc.build(elements, onFirstPage=footer, onLaterPages=footer)
